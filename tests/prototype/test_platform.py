from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
import secrets
import sqlite3
from uuid import uuid4

from fastapi.testclient import TestClient
import pytest

from app.main import create_integrated_app
from app.prototype.main import create_app
from app.prototype.ownership import assign_legacy_owner
from app.prototype.platform import Platform
from app.prototype.store import Store
from conftest import event


BASE = '/api/prototype'


def register(client, email='owner@example.test'):
    response = client.post(f'{BASE}/auth/register', json={'email': email, 'password': 'test-password-123', 'name': 'Workspace owner'})
    assert response.status_code == 201, response.text
    return response.json()


def account_headers(account, workspace=None):
    return {'Authorization': f'Bearer {account["access_token"]}', **({'X-Workspace-ID': workspace['id']} if workspace else {})}


def guest(client, workspace):
    headers = {'X-Workspace-ID': workspace['id'], 'X-Workspace-Key': secrets.token_urlsafe(32)}
    response = client.post(f'{BASE}/workspaces/guest', json={'workspace': workspace}, headers=headers)
    assert response.status_code == 201, response.text
    return headers


def test_secret_required_for_every_analytical_resource_and_guest_isolation(tmp_path, workspace, config):
    workspace['finance'] = [event('payment', 10, '2026-09-12')]
    with TestClient(create_app(tmp_path / 'platform.sqlite3', start_worker=False)) as client:
        headers = guest(client, workspace)
        definition = client.post(f'{BASE}/definitions', headers=headers, json={'name': 'cash', 'kind': 'simulation', 'config': config}).json()
        run = client.post(f'{BASE}/definitions/{definition["id"]}/runs', headers=headers, json={'snapshot': workspace, 'idempotency_key': 'run'}).json()
        routes = ['/definitions', f'/definitions/{definition["id"]}', '/runs', f'/runs/{run["id"]}', *[f'/runs/{run["id"]}/{resource}' for resource in ('results', 'series', 'events', 'scene-manifest', 'transitions')]]
        for route in routes:
            assert client.get(BASE + route, headers={'X-Workspace-ID': workspace['id']}).status_code == 401
            assert client.get(BASE + route, headers={**headers, 'X-Workspace-Key': secrets.token_urlsafe(32)}).status_code == 404
        other = deepcopy(workspace)
        other.update({'id': str(uuid4()), 'mode': 'demo'})
        other_headers = guest(client, other)
        assert client.get(f'{BASE}/runs/{run["id"]}', headers=other_headers).status_code == 404
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=other_headers).status_code == 404
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()['workspace']['mode'] == 'business'


def test_workspace_durable_open_fields_conflicts_and_no_implicit_history_rewrite(tmp_path, workspace, config):
    path = tmp_path / 'durable.sqlite3'
    workspace['finance'] = [event('payment', 10, '2026-09-12')]
    workspace['inventoryHistory'] = [{'id': 'history1', 'quantity': 5}]
    with TestClient(create_app(path, start_worker=False)) as client:
        headers = guest(client, workspace)
        definition = client.post(f'{BASE}/definitions', headers=headers, json={'name': 'cash', 'kind': 'simulation', 'config': config}).json()
        run = client.post(f'{BASE}/definitions/{definition["id"]}/runs', headers=headers, json={'snapshot': workspace, 'idempotency_key': 'original'}).json()
        changed = deepcopy(workspace)
        changed['cash']['amount'] = 250
        changed['paymentTerms'] = [{'id': 'terms1', 'days': 30}]
        changed['revision'] = 99
        payload = {'workspace': changed, 'expected_revision': 0}
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: client.put(f'{BASE}/workspaces/{workspace["id"]}', headers=headers, json=payload), range(2)))
        assert sorted(response.status_code for response in responses) == [200, 409]
        current = client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()
        assert current['revision'] == current['workspace']['revision'] == 1
        assert current['workspace']['paymentTerms'][0]['days'] == 30
        assert client.get(f'{BASE}/runs/{run["id"]}', headers=headers).json()['snapshot']['cash']['amount'] == 100
    with TestClient(create_app(path, start_worker=False)) as client:
        result = client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()
        assert result['workspace']['cash']['amount'] == 250
        assert result['workspace']['inventoryHistory'] == workspace['inventoryHistory']


def test_register_login_logout_claim_and_account_restore(tmp_path, workspace):
    path = tmp_path / 'accounts.sqlite3'
    with TestClient(create_app(path, start_worker=False)) as client:
        guest_headers = guest(client, workspace)
        user = register(client)
        assert set(user['user']) == {'id', 'email', 'name', 'created_at'}
        assert client.post(f'{BASE}/auth/login', json={'email': 'owner@example.test', 'password': 'wrong-password'}).status_code == 401
        login = client.post(f'{BASE}/auth/login', json={'email': 'OWNER@example.test', 'password': 'test-password-123'}).json()
        headers = {**guest_headers, **account_headers(login)}
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).status_code == 200
        claimed = client.post(f'{BASE}/workspaces/{workspace["id"]}/claim', headers=headers)
        assert claimed.status_code == 200
        assert claimed.json()['role'] == 'owner'
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=guest_headers).status_code == 404
        assert client.get(f'{BASE}/workspaces', headers=account_headers(login)).json()[0]['id'] == workspace['id']
        assert client.post(f'{BASE}/auth/logout', headers=account_headers(user)).status_code == 204
        assert client.get(f'{BASE}/auth/me', headers=account_headers(user)).status_code == 401
        with sqlite3.connect(path) as db:
            stored_password = db.execute('SELECT password_hash FROM platform_users').fetchone()[0]
            assert 'test-password' not in stored_password
            assert db.execute('SELECT token_hash FROM platform_sessions').fetchone()[0] != login['access_token']
    with TestClient(create_app(path, start_worker=False)) as client:
        assert client.get(f'{BASE}/auth/me', headers=account_headers(login)).json()['email'] == 'owner@example.test'
        assert client.get(f'{BASE}/workspaces', headers=account_headers(login)).json()[0]['role'] == 'owner'


@pytest.mark.parametrize('role,field,allowed', [('finance', 'cash', True), ('finance', 'stock', False), ('inventory', 'stock', True), ('inventory', 'cash', False), ('buyer', 'purchases', True), ('buyer', 'stock', False), ('viewer', 'cash', False)])
def test_actual_role_enforcement_on_workspace_and_analytical_writes(tmp_path, workspace, config, role, field, allowed):
    with TestClient(create_app(tmp_path / 'roles.sqlite3', start_worker=False)) as client:
        owner = register(client)
        member = register(client, 'member@example.test')
        owner_headers = account_headers(owner, workspace)
        assert client.post(f'{BASE}/workspaces', headers=owner_headers, json={'workspace': workspace}).status_code == 201
        response = client.post(f'{BASE}/workspaces/{workspace["id"]}/members', headers=owner_headers, json={'email': 'member@example.test', 'role': role})
        assert response.status_code == 200
        assert response.json()[0]['user_id'] == member['user']['id']
        member_headers = account_headers(member, workspace)
        changed = deepcopy(workspace)
        if field == 'cash':
            changed['cash']['amount'] = 500
        elif field == 'stock':
            changed['stock'][0]['onHand'] = 500
        else:
            changed['purchases'].append({'id': 'new-po'})
        status = client.put(f'{BASE}/workspaces/{workspace["id"]}', headers=member_headers, json={'workspace': changed, 'expected_revision': 0}).status_code
        assert status == (200 if allowed else 403)
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=member_headers).status_code == 200
        assert client.post(f'{BASE}/workspaces/{workspace["id"]}/members', headers=member_headers, json={'email': 'owner@example.test', 'role': 'viewer'}).status_code == 403
        cash = client.post(f'{BASE}/definitions', headers=member_headers, json={'name': 'cash', 'kind': 'simulation', 'config': config})
        forecast = client.post(f'{BASE}/definitions', headers=member_headers, json={'name': 'forecast', 'kind': 'forecast', 'config': {**config, 'product_id': 'p1'}})
        assert cash.status_code == (201 if role == 'finance' else 403)
        assert forecast.status_code == (201 if role in {'inventory', 'buyer'} else 403)


def test_last_owner_membership_revocation_and_archive_preserve_records(tmp_path, workspace):
    with TestClient(create_app(tmp_path / 'ownership.sqlite3', start_worker=False)) as client:
        owner = register(client)
        member = register(client, 'member@example.test')
        headers = account_headers(owner, workspace)
        client.post(f'{BASE}/workspaces', headers=headers, json={'workspace': workspace})
        member_path = f'{BASE}/workspaces/{workspace["id"]}/members'
        assert client.delete(f'{member_path}/{owner["user"]["id"]}', headers=headers).status_code == 409
        assert client.post(member_path, headers=headers, json={'email': 'unregistered@example.test', 'role': 'viewer'}).status_code == 404
        assert client.post(member_path, headers=headers, json={'email': member['user']['email'], 'role': 'viewer'}).status_code == 200
        member_headers = account_headers(member, workspace)
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=member_headers).status_code == 200
        assert client.delete(f'{member_path}/{member["user"]["id"]}', headers=headers).status_code == 200
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=member_headers).status_code == 404
        assert client.delete(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()['archived']
        assert client.get(f'{BASE}/workspaces', headers=headers).json() == []
        assert client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).status_code == 200
        assert client.put(f'{BASE}/workspaces/{workspace["id"]}', headers=headers, json={'workspace': workspace, 'expected_revision': 0}).status_code == 403
        assert client.post(f'{BASE}/workspaces/{workspace["id"]}/restore', headers=headers).json()['archived'] is False


def test_native_xlsx_multiple_sheets_preview_does_not_commit(tmp_path, workspace):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = 'Instructions'
    book.active.append(['Description'])
    book.active.append(['Select Sales'])
    sales = book.create_sheet('Sales')
    sales.append(['Date', 'SKU', 'Quantity', 'Amount'])
    sales.append([datetime(2026, 9, 11), 'ABC', 2, 24.5])
    target = BytesIO()
    book.save(target)
    with TestClient(create_app(tmp_path / 'excel.sqlite3', start_worker=False)) as client:
        headers = guest(client, workspace)
        response = client.post(f'{BASE}/imports/preview', headers=headers, files={'file': ('sales.xlsx', target.getvalue())}, data={'sheet_name': 'Sales'})
        assert response.status_code == 200, response.text
        preview = response.json()
        assert preview['sheets'] == [{'name': 'Instructions', 'rowCount': 1}, {'name': 'Sales', 'rowCount': 1}]
        assert preview['headers'] == ['Date', 'SKU', 'Quantity', 'Amount']
        assert preview['rows'] == [['2026-09-11', 'ABC', '2', '24.5']]
        assert preview['selectedSheet'] == 'Sales'
        again = client.post(f'{BASE}/imports/preview', headers=headers, files={'file': ('sales.xlsx', target.getvalue())}, data={'sheet_name': 'Sales'}).json()
        assert again['fingerprint'] == preview['fingerprint']
        current = client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()
        assert current['revision'] == 0 and current['workspace']['sales'] == workspace['sales']
        assert client.post(f'{BASE}/imports/preview', headers=headers, files={'file': ('bad.xlsx', b'not a workbook')}).status_code == 422
        assert client.post(f'{BASE}/imports/preview', headers=headers, files={'file': ('sales.xlsx', target.getvalue())}, data={'sheet_name': 'Unknown'}).status_code == 422
        assert client.post(f'{BASE}/imports/preview', headers=headers, files={'file': ('huge.xlsx', b'x' * 5_000_001)}).status_code == 413


def test_native_xls_preview_reads_numeric_and_date_cells(tmp_path, workspace):
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('Sales')
    for column, name in enumerate(['Date', 'SKU', 'Quantity']):
        sheet.write(0, column, name)
    sheet.write(1, 0, datetime(2026, 9, 11), xlwt.easyxf(num_format_str='YYYY-MM-DD'))
    sheet.write(1, 1, 'LEGACY')
    sheet.write(1, 2, 3)
    output = BytesIO()
    book.save(output)
    with TestClient(create_app(tmp_path / 'xls.sqlite3', start_worker=False)) as client:
        headers = guest(client, workspace)
        response = client.post(f'{BASE}/workspaces/{workspace["id"]}/imports/preview', headers=headers, files={'file': ('legacy.xls', output.getvalue())})
        assert response.status_code == 200, response.text
        assert response.json()['rows'] == [['2026-09-11', 'LEGACY', '3']]


def test_existing_uuid_cannot_be_claimed_over_http_and_cli_migration_is_explicit(tmp_path, workspace, config):
    path = tmp_path / 'legacy.sqlite3'
    store = Store(path)
    workspace['finance'] = [event('payment', 10, '2026-09-12')]
    definition = store.create_definition(workspace['id'], {'name': 'Legacy', 'kind': 'simulation', 'config': config})
    run = store.submit(workspace['id'], definition['id'], {'snapshot': workspace, 'idempotency_key': 'legacy'})
    with TestClient(create_app(path, start_worker=False)) as client:
        user = register(client)
        key = {'X-Workspace-Key': secrets.token_urlsafe(32), 'X-Workspace-ID': workspace['id']}
        assert client.post(f'{BASE}/workspaces/guest', headers=key, json={'workspace': workspace}).status_code == 409
        assert client.post(f'{BASE}/workspaces/{workspace["id"]}/claim', headers={**key, **account_headers(user)}).status_code == 404
        assign_legacy_owner(str(path), workspace['id'], user['user']['email'])
        headers = account_headers(user, workspace)
        restored = client.get(f'{BASE}/runs/{run["id"]}', headers=headers).json()
        assert restored['snapshot'] == workspace
        assert client.get(f'{BASE}/workspaces', headers=headers).json()[0]['id'] == workspace['id']


def test_original_app_has_platform_lifespan_auth_alias_and_no_postgres_requirement(tmp_path, workspace):
    with TestClient(create_integrated_app(tmp_path / 'main.sqlite3', start_worker=False)) as client:
        auth = client.post('/api/v1/auth/register', json={'email': 'main@example.test', 'password': 'test-password-123'})
        assert auth.status_code == 201
        headers = account_headers(auth.json(), workspace)
        assert client.post(f'{BASE}/workspaces', headers=headers, json={'workspace': workspace}).status_code == 201
        assert client.get('/api/v1/auth/me', headers=headers).json()['email'] == 'main@example.test'
        assert client.get(f'{BASE}/health').json()['persistence'] == 'sqlite'


def test_private_keys_and_session_expiry_are_enforced(tmp_path, workspace):
    path = tmp_path / 'expiry.sqlite3'
    with TestClient(create_app(path, start_worker=False)) as client:
        assert client.post(f'{BASE}/workspaces/guest', json={'workspace': workspace}).status_code == 422
        headers = guest(client, workspace)
        assert client.post(f'{BASE}/workspaces/guest', headers=headers, json={'workspace': workspace}).status_code == 201
        user = register(client)
        with sqlite3.connect(path) as db:
            assert db.execute('SELECT guest_key_hash FROM platform_workspaces').fetchone()[0] != headers['X-Workspace-Key']
            db.execute("UPDATE platform_sessions SET expires_at='2000-01-01'")
        assert client.get(f'{BASE}/auth/me', headers=account_headers(user)).status_code == 401


def test_export_and_consistent_backup_retain_current_records_accounts_and_analytical_history(tmp_path, workspace, config):
    from app.prototype.backup import backup_database
    path = tmp_path / 'original.sqlite3'
    backup = tmp_path / 'backup.sqlite3'
    workspace['finance'] = [event('payment', 10, '2026-09-12')]
    with TestClient(create_app(path, start_worker=False)) as client:
        headers = guest(client, workspace)
        export = client.get(f'{BASE}/workspaces/{workspace["id"]}/export', headers=headers)
        assert export.json() == workspace
        assert export.headers['content-disposition'].startswith('attachment;')
        saved = client.post(f'{BASE}/definitions', headers=headers, json={'name': 'Retained cash', 'kind': 'simulation', 'config': config}).json()
        run = client.post(f'{BASE}/definitions/{saved["id"]}/runs', headers=headers, json={'snapshot': workspace, 'idempotency_key': 'backup'}).json()
        backup_database(str(path), str(backup))
        with pytest.raises(ValueError, match='never overwritten'):
            backup_database(str(path), str(backup))
    with TestClient(create_app(backup, start_worker=False)) as restored:
        assert restored.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()['workspace'] == workspace
        assert restored.get(f'{BASE}/runs/{run["id"]}', headers=headers).json()['snapshot'] == workspace


def test_login_rate_limit_is_persisted_and_generic(tmp_path):
    path = tmp_path / 'rate.sqlite3'
    with TestClient(create_app(path, start_worker=False)) as client:
        register(client)
        for _ in range(10):
            assert client.post(f'{BASE}/auth/login', json={'email': 'owner@example.test', 'password': 'incorrect-password'}).status_code == 401
        assert client.post(f'{BASE}/auth/login', json={'email': 'owner@example.test', 'password': 'incorrect-password'}).status_code == 429
    with TestClient(create_app(path, start_worker=False)) as client:
        assert client.post(f'{BASE}/auth/login', json={'email': 'owner@example.test', 'password': 'incorrect-password'}).status_code == 429


def test_finance_role_can_preserve_pending_records_without_converting_unknowns_to_zero(tmp_path, workspace):
    path = tmp_path / 'pending.sqlite3'
    with TestClient(create_app(path, start_worker=False)) as client:
        owner = register(client)
        member = register(client, 'finance@example.test')
        owner_headers = account_headers(owner, workspace)
        client.post(f'{BASE}/workspaces', headers=owner_headers, json={'workspace': workspace})
        client.post(f'{BASE}/workspaces/{workspace["id"]}/members', headers=owner_headers, json={'email': 'finance@example.test', 'role': 'finance'})
        changed = deepcopy(workspace)
        changed['pendingFinance'] = [{'id': 'incomplete-bill', 'sourceId': 'manual-pending', 'kind': 'payable', 'name': 'Invoice amount pending', 'counterparty': '', 'currency': 'MXN', 'amount': None, 'paidAmount': 0, 'dueDate': None, 'expectedDate': None}]
        headers = account_headers(member, workspace)
        response = client.put(f'{BASE}/workspaces/{workspace["id"]}', headers=headers, json={'workspace': changed, 'expected_revision': 0})
        assert response.status_code == 200, response.text
        assert response.json()['workspace']['finance'] == workspace['finance']
    with TestClient(create_app(path, start_worker=False)) as client:
        restored = client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()['workspace']
        assert restored['pendingFinance'][0]['amount'] is None
        assert restored['pendingFinance'][0]['paidAmount'] == 0


def test_dated_finance_observations_persist_under_finance_role_without_changing_cash(tmp_path, workspace):
    workspace['finance'] = [event('invoice', 100, '2026-09-12', 'receivable', paidAmount=40)]
    with TestClient(create_app(tmp_path / 'history.sqlite3', start_worker=False)) as client:
        owner = register(client)
        member = register(client, 'finance-history@example.test')
        headers = account_headers(owner, workspace)
        client.post(f'{BASE}/workspaces', headers=headers, json={'workspace': workspace})
        client.post(f'{BASE}/workspaces/{workspace["id"]}/members', headers=headers, json={'email': 'finance-history@example.test', 'role': 'finance'})
        changed = deepcopy(workspace)
        changed['financeEvents'] = [{'id': 'received', 'sourceId': 'manual-receipt', 'kind': 'customer_collection', 'recordId': 'invoice', 'paymentReference': 'BANK-123', 'date': '2026-09-03', 'amount': 30, 'currency': 'MXN'}]
        response = client.put(f'{BASE}/workspaces/{workspace["id"]}', headers=account_headers(member, workspace), json={'workspace': changed, 'expected_revision': 0})
        assert response.status_code == 200, response.text
        result = client.get(f'{BASE}/workspaces/{workspace["id"]}', headers=headers).json()['workspace']
        assert result['financeEvents'] == changed['financeEvents']
        assert result['cash'] == workspace['cash']
        assert result['finance'] == workspace['finance']
