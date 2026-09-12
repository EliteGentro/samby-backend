from datetime import date, datetime, time
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from .store import ResourceError


MAX_ROWS = 10_000
MAX_COLUMNS = 200


def display(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def preview(filename: str, content: bytes, sheet_name: str | None) -> dict:
    extension = Path(filename).suffix.lower()
    if extension not in {'.xlsx', '.xls'}:
        raise ResourceError(422, 'Choose a native .xlsx or .xls workbook. CSV continues through its reviewed CSV workflow.')
    if not content or len(content) > 5_000_000:
        raise ResourceError(413, 'Choose a nonempty workbook no larger than 5 MB.')
    warnings = []
    try:
        if extension == '.xlsx':
            from openpyxl import load_workbook
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 50_000_000:
                    raise ResourceError(413, 'The expanded workbook exceeds the safe preview size. Export a smaller sheet.')
            book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            sheets = [{'name': sheet.title, 'rowCount': max(0, (sheet.max_row or 0) - 1)} for sheet in book.worksheets]
            names = [sheet['name'] for sheet in sheets]
            selected = sheet_name or (names[0] if names else None)
            if selected not in names:
                book.close()
                raise ResourceError(422, 'Select a worksheet present in the workbook.')
            sheet = book[selected]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                book.close()
                raise ResourceError(413, 'Select a sheet with at most 10,000 data rows and 200 columns.')
            values = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index > MAX_ROWS or len(row) > MAX_COLUMNS:
                    book.close()
                    raise ResourceError(413, 'The sheet exceeds 10,000 data rows or 200 columns.')
                values.append([display(value) for value in row])
            book.close()
            warnings.append('Formula cells use workbook-cached values only. Missing cached values remain blank; formulas and external links are not executed.')
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            names = book.sheet_names()
            sheets = [{'name': sheet.name, 'rowCount': max(0, sheet.nrows - 1)} for sheet in book.sheets()]
            selected = sheet_name or (names[0] if names else None)
            if selected not in names:
                book.release_resources()
                raise ResourceError(422, 'Select a worksheet present in the workbook.')
            sheet = book.sheet_by_name(selected)
            if sheet.nrows > MAX_ROWS + 1 or sheet.ncols > MAX_COLUMNS:
                book.release_resources()
                raise ResourceError(413, 'Select a sheet with at most 10,000 data rows and 200 columns.')
            values = []
            for row in sheet.get_rows():
                converted = []
                for cell in row:
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#ERROR')
                    converted.append(display(value))
                values.append(converted)
            book.release_resources()
            warnings.append('Legacy XLS formulas use saved workbook results only; macros and formulas are not executed.')
    except ResourceError:
        raise
    except Exception as error:
        raise ResourceError(422, 'This workbook could not be read. Remove encryption or repair it in Excel, then save a valid .xlsx or .xls file.') from error
    while values and not any(value.strip() for value in values[0]):
        values.pop(0)
    while values and not any(value.strip() for value in values[-1]):
        values.pop()
    if len(values) < 2:
        raise ResourceError(422, 'The selected sheet needs a header row and at least one data row.')
    width = max((max((i + 1 for i, value in enumerate(row) if value.strip()), default=0) for row in values), default=0)
    headers = [value.strip() for value in values[0][:width]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ResourceError(422, 'Column headings must be nonempty and unique. Review the header row before mapping.')
    rows = [(row + [''] * width)[:width] for row in values[1:]]
    warnings.append('Dates use ISO format and numeric cells use their stored numeric value; currency symbols and custom display formats do not establish amount meaning.')
    return {'sheets': sheets, 'selectedSheet': selected, 'headers': headers, 'rows': rows, 'fingerprint': hashlib.sha256(content + b'\0' + selected.encode()).hexdigest(), 'sourceName': Path(filename).name, 'warnings': warnings}
